#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


NAVY = RGBColor(22, 50, 79)
BLUE = RGBColor(47, 102, 144)
TEAL = RGBColor(58, 157, 143)
ORANGE = RGBColor(224, 122, 95)
RED = RGBColor(196, 69, 54)
GRAY = RGBColor(103, 117, 132)
LIGHT = RGBColor(234, 240, 245)
WHITE = RGBColor(255, 255, 255)
FONT = "Microsoft YaHei"


def add_text(slide, x, y, w, h, text, size=20, color=NAVY, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True; frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = FONT; paragraph.font.size = Pt(size); paragraph.font.bold = bold; paragraph.font.color.rgb = color
    return box


def add_rich_lines(slide, x, y, w, h, lines, size=17, color=NAVY, bullet=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = FONT; paragraph.font.size = Pt(size); paragraph.font.color.rgb = color
        paragraph.space_after = Pt(9)
        if bullet:
            paragraph.text = "• " + paragraph.text
    return box


def add_title(slide, title, subtitle=None):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    add_text(slide, 0.55, 0.28, 12.2, 0.56, title, size=25, bold=True)
    if subtitle:
        add_text(slide, 0.58, 0.84, 12.0, 0.38, subtitle, size=11, color=GRAY)


def add_footer(slide, number, source="本地结果审计 · 2026-08-10"):
    add_text(slide, 0.55, 7.15, 11.8, 0.22, source, size=8, color=GRAY)
    add_text(slide, 12.35, 7.12, 0.4, 0.24, str(number), size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def add_picture_contain(slide, path: Path, x, y, w, h):
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio > box_ratio:
        picture_w = w; picture_h = w / ratio
        picture_x = x; picture_y = y + (h - picture_h) / 2
    else:
        picture_h = h; picture_w = h * ratio
        picture_x = x + (w - picture_w) / 2; picture_y = y
    return slide.shapes.add_picture(str(path), Inches(picture_x), Inches(picture_y), Inches(picture_w), Inches(picture_h))


def add_metric_card(slide, x, y, w, h, value, label, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT; shape.line.color.rgb = color
    add_text(slide, x + 0.14, y + 0.14, w - 0.28, 0.55, value, size=27, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.14, y + 0.78, w - 0.28, h - 0.88, label, size=12, color=NAVY, align=PP_ALIGN.CENTER)


def build_presentation(export_dir: Path, output: Path) -> None:
    summary = json.loads((export_dir / "summary_metrics.json").read_text(encoding="utf-8"))
    old = summary["old_loco"]
    fixed = summary["minimal_fixed_blca"]
    pf = summary["patient_fold"]
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.22))
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
    add_text(slide, 0.72, 0.75, 11.9, 0.95, "CC-HHGT v3.0 state 训练效果审计", size=34, bold=True)
    add_text(slide, 0.75, 1.72, 11.8, 0.72, "sampler 已修复 · 最小 LOCO 已重训 · 250 epoch 充分性已验证", size=20, color=BLUE)
    add_metric_card(slide, 0.75, 2.75, 3.65, 1.68, "297 / 297", "旧 LOCO state 训练折全部是正样本\n旧全局指标不可解释", RED)
    add_metric_card(slide, 4.84, 2.75, 3.65, 1.68, f"{fixed['old_test_auroc']:.3f} → {fixed['fixed_test_auroc']:.3f}", "同一 RGCN / BLCA test AUROC\n修复后跨癌种信号出现", TEAL)
    add_metric_card(slide, 8.93, 2.75, 3.65, 1.68, f"{fixed['best_epoch']} / {fixed['stop_epoch']}", "best epoch / early-stop epoch\n上限 250 有余量", BLUE)
    add_text(slide, 0.85, 5.05, 11.6, 0.70, "核心结论：不是 stemness 数据结构有问题，而是 EXTEND 与全部 stemness 共用的 state PU 采样器把 unlabeled 清空了。", size=21, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(slide, 0.85, 6.04, 11.6, 0.55, "当前监督训练流程已恢复并完成本地 GPU 最小验证；旧 strict_release 的 state 结果仍需全矩阵重跑后才能替换。", size=14, color=GRAY, align=PP_ALIGN.CENTER)
    add_footer(slide, 1)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "1. 为什么 stemness 和 EXTEND 一起失败？", "最终表结构相似，不代表训练入口与采样路径相同")
    add_picture_contain(slide, export_dir / "01_sampling_bug_and_fix.png", 0.45, 1.20, 8.0, 5.55)
    add_rich_lines(
        slide, 8.62, 1.34, 4.15, 4.8,
        [
            "pathway 使用 sample_training_pairs：先抽 strong/weak positive，再按 1:4 加入 unlabeled。",
            "state 使用独立 _sample_train：旧逻辑先无条件保留所有 positive。",
            "BLCA 折 positive=182,669 > cap=150,000，导致 unlabeled 配额直接变成 0。",
            "EXTEND 与 6 个 stemness state 共用这个 sampler，所以都会失败；不是 stemness parquet 特有问题。",
            "修复后按 cancer × state 分层，生产 cap 为 P=30,000 / U=120,000，并增加单类硬报错。",
        ], size=16,
    )
    add_footer(slide, 2)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "2. 旧跨癌种结果确实无效", "完整矩阵审计：3 模型 × 33 癌种 × 3 seed")
    add_picture_contain(slide, export_dir / "02_old_loco_global_audit.png", 0.45, 1.12, 9.0, 5.72)
    add_rich_lines(
        slide, 9.58, 1.52, 3.18, 4.7,
        [
            f"旧 LOCO median AUROC = {old['median_auroc']:.3f}。",
            f"mean ± std = {old['mean_auroc']:.3f} ± {old['std_auroc']:.3f}。",
            f"AUROC ≥ 0.6 仅 {old['fraction_ge_0_6']*100:.1f}%。",
            "但这些数值首先反映全正训练退化，不能用于证明‘跨癌种本质上不可预测’。",
            "pathway 任务没有走这个 state sampler，不能由此宣布 pathway 结果同步失效。",
        ], size=15,
    )
    add_footer(slide, 3)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "3. 最小重训：修复后能学习，250 步对本实验足够", "RGCN · LOCO_BLCA · seed 20260726 · state 20,000（P:U=1:4）")
    add_picture_contain(slide, export_dir / "03_minimal_training_curves.png", 0.35, 1.08, 9.25, 5.95)
    add_metric_card(slide, 9.72, 1.35, 2.9, 1.3, f"{fixed['fixed_val_auroc']:.3f}", "validation state AUROC", TEAL)
    add_metric_card(slide, 9.72, 2.88, 2.9, 1.3, f"{fixed['fixed_test_auroc']:.3f}", "BLCA test state AUROC", BLUE)
    add_metric_card(slide, 9.72, 4.41, 2.9, 1.3, f"+{fixed['fixed_test_auroc']-fixed['old_test_auroc']:.3f}", "相对旧版 test AUROC", ORANGE)
    add_text(slide, 9.76, 5.95, 2.85, 0.76, "epoch 132 最佳；157 自动早停。\n80/140 上限均不足，250 有余量。", size=13, color=NAVY, align=PP_ALIGN.CENTER)
    add_footer(slide, 4)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "4. stemness 逐 state 结果：修复有用，但难度仍高", "同一 BLCA test；pooled AUROC 会被不同 state 的阳性率差异抬高")
    add_picture_contain(slide, export_dir / "04_stemness_extend_old_vs_fixed.png", 0.38, 1.12, 9.3, 5.78)
    add_rich_lines(
        slide, 9.78, 1.38, 2.92, 4.9,
        [
            f"pooled AUROC：{fixed['old_test_auroc']:.3f} → {fixed['fixed_test_auroc']:.3f}。",
            f"7-state macro AUROC：{fixed['macro_state_old_auroc']:.3f} → {fixed['macro_state_fixed_auroc']:.3f}。",
            "RNAss 提升最清楚；EREG.EXPss 仍接近随机。",
            "因此修复解决了训练退化，但尚未证明每种 stemness 都达到可交付性能。",
            "后续主指标应报告 state 内 macro AUROC，而不是只报告 pooled AUROC。",
        ], size=14,
    )
    add_footer(slide, 5)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "5. PF 的 0.72 是关联重现性，不是图模型净增益", "同癌种内用训练患者估计效应，再测试 held-out 患者子集是否重现")
    add_picture_contain(slide, export_dir / "05_patient_fold_model_vs_baseline.png", 0.38, 1.15, 8.8, 5.7)
    add_rich_lines(
        slide, 9.32, 1.35, 3.45, 5.2,
        [
            f"495 runs，{pf['valid_runs']} 个 AUROC 有效。",
            f"4-head ensemble median = {pf['model_median_auroc']:.3f}。",
            f"简单 abs(train_effect) baseline median = {pf['baseline_median_auroc']:.3f}。",
            f"ensemble − baseline：median {pf['median_model_minus_baseline']:+.3f}，mean {pf['mean_model_minus_baseline']:+.3f}。",
            f"按 state 内 macro 后 median 降到 {pf['macro_state_median_auroc']:.3f}。",
            "结论：同癌种关联可重现，但当前图 ensemble 没有超过直接效应量；不能把 0.72 归因于图学习。",
        ], size=14,
    )
    add_footer(slide, 6)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "6. 当前状态与下一步")
    add_text(slide, 0.75, 1.15, 5.8, 0.42, "已完成", size=22, bold=True, color=TEAL)
    add_rich_lines(
        slide, 0.78, 1.68, 5.75, 3.75,
        [
            "修复 state PU sampler：1:4 P/U、cancer×state 分层、单类 fail-closed。",
            "修复 pathway 大表内存路径：逐癌种统计并采样，不再拼接 3,293 万行。",
            "12 项相关回归测试通过；RTX 4070 Ti SUPER 完成 250/25 最小训练。",
            "生成逐 epoch、逐 state、PF 基线对比的可复核图表与 TSV。",
        ], size=16,
    )
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.78), Inches(5.47), Inches(5.67), Inches(0.95))
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT; shape.line.color.rgb = TEAL
    add_text(slide, 0.98, 5.65, 5.28, 0.56, "监督训练没有停；最小验证已完整跑通。", size=19, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    add_text(slide, 6.95, 1.15, 5.6, 0.42, "仍需完成", size=22, bold=True, color=ORANGE)
    add_rich_lines(
        slide, 6.98, 1.68, 5.55, 4.65,
        [
            "旧 strict_release 的 297 个 state LOCO 结果全部作废，不能继续用于发布结论。",
            "按修复代码重跑 3 模型 × 33 癌种 × 3 seed；本次单折结果只证明 bug 修复有效，不代表全癌种性能。",
            "新增发布门槛：每折训练 P/U 审计、pooled + state-macro 双指标、per-state 最小表现。",
            "PF 报告必须同时给 abs(train_effect) baseline；若图模型不超过基线，应调整模型而不是只调 epoch。",
            "生产建议保留 epochs=250 / patience=25；训练步数不是当前首要瓶颈。",
        ], size=16,
    )
    add_footer(slide, 7, source=f"产物目录：{export_dir}")

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--export-dir", default=r"D:\model\exports\state_training_effect_fixed")
    parser.add_argument("--output", default=r"D:\model\exports\state_training_effect_fixed\CC_HHGT_v3_state_training_effect_report.pptx")
    args = parser.parse_args()
    build_presentation(Path(args.export_dir).resolve(), Path(args.output).resolve())
    print(Path(args.output).resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
