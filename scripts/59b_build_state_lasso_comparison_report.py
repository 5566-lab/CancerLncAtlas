#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path
import textwrap

import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


RNASS = "stemness_rna::RNAss"
DNASS = "stemness_dna::DNAss"
TARGETS = [RNASS, DNASS]
TARGET_LABELS = {RNASS: "RNAss", DNASS: "DNAss"}
MODEL_LABELS = {
    "rgcn": "RGCN",
    "hgt": "HGT",
    "cc_hhgt_strict": "CC-HHGT",
    "rgcn_hgt_equal": "RGCN+HGT",
    "three_model_equal": "3-model equal",
    "lasso_abs_coefficient": "LASSO coefficient",
    "train_abs_effect": "Train effect",
    "patient_head": "Patient head",
    "graph_heads_equal": "Graph heads",
    "all_heads_equal": "All heads",
}
COLORS = {
    "navy": "#15324A",
    "blue": "#2F6690",
    "teal": "#3A9D8F",
    "orange": "#E07A5F",
    "gold": "#E9C46A",
    "gray": "#7A8491",
    "light": "#EAF0F5",
    "red": "#C44536",
}


def configure_plotting() -> None:
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    if font_path.exists():
        font_manager.fontManager.addfont(str(font_path))
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font_path)).get_name()
    plt.rcParams.update(
        {
            "font.size": 11,
            "axes.titlesize": 14,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )


def require_pass(path: Path, name: str) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("status") != "PASS":
        raise RuntimeError(f"{name} audit is not PASS: {path}")
    return payload


def save_loco_plot(oof_summary: pd.DataFrame, output: Path) -> None:
    models = ["rgcn", "hgt", "cc_hhgt_strict", "rgcn_hgt_equal", "three_model_equal"]
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.5), sharey=True)
    colors = [COLORS["blue"], COLORS["teal"], COLORS["orange"], COLORS["gold"], COLORS["navy"]]
    for axis, state_id in zip(axes, TARGETS, strict=True):
        state = oof_summary.loc[oof_summary.state_id.eq(state_id)].set_index("model_name")
        values = [float(state.loc[model, "pooled_auroc"]) for model in models]
        bars = axis.bar(np.arange(len(models)), values, color=colors)
        axis.axhline(0.5, color=COLORS["gray"], linestyle="--", linewidth=1.3)
        axis.set_xticks(np.arange(len(models)), [MODEL_LABELS[model] for model in models], rotation=24, ha="right")
        axis.set_ylim(0.42, max(0.82, max(values) + 0.08))
        axis.set_title(f"LOCO pooled AUROC — {TARGET_LABELS[state_id]}")
        for bar, value in zip(bars, values, strict=True):
            axis.text(bar.get_x() + bar.get_width() / 2, value + 0.012, f"{value:.3f}", ha="center", fontsize=10)
    axes[0].set_ylabel("AUROC（完全未见癌种）")
    fig.tight_layout()
    fig.savefig(output, dpi=190, bbox_inches="tight")
    plt.close(fig)


def save_pf_plot(summary: pd.DataFrame, output: Path) -> None:
    models = ["lasso_abs_coefficient", "graph_heads_equal", "patient_head", "all_heads_equal"]
    x = np.arange(len(TARGETS))
    width = 0.19
    fig, axis = plt.subplots(figsize=(11.8, 5.8))
    colors = [COLORS["gray"], COLORS["teal"], COLORS["orange"], COLORS["navy"]]
    for index, (model, color) in enumerate(zip(models, colors, strict=True)):
        model_frame = summary.loc[summary.model_name.eq(model)].set_index("state_id")
        values = [float(model_frame.loc[state, "median_auroc"]) for state in TARGETS]
        positions = x + (index - 1.5) * width
        bars = axis.bar(positions, values, width=width, color=color, label=MODEL_LABELS[model])
        for bar, value in zip(bars, values, strict=True):
            axis.text(bar.get_x() + bar.get_width() / 2, value + 0.012, f"{value:.3f}", ha="center", fontsize=9)
    axis.axhline(0.5, color=COLORS["gray"], linestyle="--", linewidth=1.2)
    axis.set_xticks(x, [TARGET_LABELS[state] for state in TARGETS])
    axis.set_ylim(0.45, 0.86)
    axis.set_ylabel("患者折内关联重现 AUROC（中位数）")
    axis.set_title("同一患者折、同一 held-out association label 的匹配比较")
    axis.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(output, dpi=190, bbox_inches="tight")
    plt.close(fig)


def save_endpoint_plot(lasso_patient: pd.DataFrame, association: pd.DataFrame, output: Path) -> None:
    patient = lasso_patient.set_index("state_id")
    coefficient = association.loc[association.model_name.eq("lasso_abs_coefficient")].set_index("state_id")
    score_auc = [float(patient.loc[state, "median_test_high_low_auroc"]) for state in TARGETS]
    association_auc = [float(coefficient.loc[state, "median_auroc"]) for state in TARGETS]
    x = np.arange(len(TARGETS))
    fig, axis = plt.subplots(figsize=(10.8, 5.7))
    first = axis.bar(x - 0.18, score_auc, 0.36, label="患者高/低 stemness 预测", color=COLORS["blue"])
    second = axis.bar(x + 0.18, association_auc, 0.36, label="LASSO 系数预测关联重现", color=COLORS["gray"])
    axis.axhline(0.5, color=COLORS["gray"], linestyle="--", linewidth=1.2)
    axis.set_xticks(x, [TARGET_LABELS[state] for state in TARGETS])
    axis.set_ylim(0.45, 1.02)
    axis.set_ylabel("AUROC 中位数")
    axis.set_title("同一个 LASSO，在两个不同任务上的结论完全不同")
    axis.legend(frameon=False)
    for bars in (first, second):
        for bar in bars:
            axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.012, f"{bar.get_height():.3f}", ha="center")
    fig.tight_layout()
    fig.savefig(output, dpi=190, bbox_inches="tight")
    plt.close(fig)


def save_delta_plot(tests: pd.DataFrame, output: Path) -> None:
    models = ["graph_heads_equal", "patient_head", "all_heads_equal"]
    x = np.arange(len(TARGETS))
    width = 0.24
    fig, axis = plt.subplots(figsize=(10.8, 5.7))
    for index, (model, color) in enumerate(zip(models, [COLORS["teal"], COLORS["orange"], COLORS["navy"]], strict=True)):
        state = tests.loc[tests.model_name.eq(model)].set_index("state_id")
        values = [float(state.loc[target, "median_delta_auroc_vs_lasso"]) for target in TARGETS]
        bars = axis.bar(x + (index - 1) * width, values, width=width, label=MODEL_LABELS[model], color=color)
        for bar, value in zip(bars, values, strict=True):
            axis.text(bar.get_x() + bar.get_width() / 2, value + (0.008 if value >= 0 else -0.018), f"{value:+.3f}", ha="center", fontsize=9)
    axis.axhline(0.0, color=COLORS["gray"], linewidth=1.2)
    axis.set_xticks(x, [TARGET_LABELS[state] for state in TARGETS])
    axis.set_ylabel("配对 ΔAUROC vs LASSO（中位数）")
    axis.set_title("关联重现任务：每个模型与 LASSO 在相同 fold/candidates 上配对")
    axis.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(output, dpi=190, bbox_inches="tight")
    plt.close(fig)


def save_cancer_heatmap(cancer_metrics: pd.DataFrame, output: Path) -> None:
    frame = cancer_metrics.loc[
        cancer_metrics.model_name.eq("three_model_equal") & cancer_metrics.state_id.isin(TARGETS)
    ].copy()
    pivot = frame.pivot(index="cancer_id", columns="state_id", values="auroc").reindex(columns=TARGETS)
    fig, axis = plt.subplots(figsize=(6.8, 10.5))
    image = axis.imshow(pivot.to_numpy(float), aspect="auto", vmin=0.35, vmax=0.80, cmap="RdYlBu")
    axis.set_xticks(np.arange(2), [TARGET_LABELS[state] for state in TARGETS])
    axis.set_yticks(np.arange(len(pivot)), pivot.index.tolist(), fontsize=8)
    axis.set_title("3-model equal：各留出癌种 AUROC")
    for row in range(len(pivot)):
        for column in range(2):
            value = pivot.iloc[row, column]
            if np.isfinite(value):
                axis.text(column, row, f"{value:.2f}", ha="center", va="center", fontsize=7)
    fig.colorbar(image, ax=axis, shrink=0.70, label="AUROC")
    fig.tight_layout()
    fig.savefig(output, dpi=190, bbox_inches="tight")
    plt.close(fig)


def set_text_style(paragraph, size: float, bold: bool = False, color: str = COLORS["navy"]) -> None:
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = bytes.fromhex(color.removeprefix("#"))


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.82))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    set_text_style(paragraph, 25, True)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.88), Inches(12.0), Inches(0.38))
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        set_text_style(p, 10.5, False, COLORS["gray"])


def add_text(slide, text: str, left: float, top: float, width: float, height: float, size: float = 16, bold: bool = False, color: str = COLORS["navy"]) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    set_text_style(paragraph, size, bold, color)


def add_bullets(slide, items: list[str], left: float, top: float, width: float, height: float, size: float = 16) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.space_after = Pt(9)
        set_text_style(paragraph, size)


def add_picture(slide, path: Path, left: float, top: float, width: float) -> None:
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def build_markdown(
    matrix_audit: dict,
    task_metrics: pd.DataFrame,
    oof_summary: pd.DataFrame,
    matched_summary: pd.DataFrame,
    lasso_patient: pd.DataFrame,
    paired_tests: pd.DataFrame,
    output: Path,
) -> None:
    quality = task_metrics.state_training_positive_rate.astype(float)
    selected_oof = oof_summary.loc[
        oof_summary.state_id.isin(TARGETS) & oof_summary.model_name.isin(["rgcn", "hgt", "cc_hhgt_strict", "three_model_equal"]),
        ["model_name", "state_id", "pooled_auroc", "median_cancer_auroc", "pooled_auprc"],
    ].copy()
    selected_oof["model_name"] = selected_oof.model_name.map(MODEL_LABELS)
    selected_oof["state_id"] = selected_oof.state_id.map(TARGET_LABELS)
    selected_pf = matched_summary.loc[
        matched_summary.state_id.isin(TARGETS) & matched_summary.model_name.isin(["lasso_abs_coefficient", "graph_heads_equal", "patient_head", "all_heads_equal"]),
        ["model_name", "state_id", "n_valid_fold_units", "median_auroc", "mean_auroc", "median_auprc"],
    ].copy()
    selected_pf["model_name"] = selected_pf.model_name.map(MODEL_LABELS)
    selected_pf["state_id"] = selected_pf.state_id.map(TARGET_LABELS)
    patient = lasso_patient.copy()
    patient["state_id"] = patient.state_id.map(TARGET_LABELS)
    lines = [
        "# CC-HHGT V3.0 RNAss/DNAss 完整重训与 LASSO 对照",
        "",
        "## 审计结论",
        "",
        f"- 完整矩阵：{matrix_audit['valid_task_metrics']}/{matrix_audit['expected_tasks']} 个任务通过。",
        f"- state 训练正样本率：{quality.min():.3f}–{quality.max():.3f}；每任务均为 150,000 条、P:U=1:4。",
        "- 所有正式数字来自 isolated fixed-sampler 输出；旧版全正样本矩阵不参与结论。",
        "",
        "## 跨癌 LOCO（完全未见癌种）",
        "",
        selected_oof.to_markdown(index=False, floatfmt=".4f"),
        "",
        "## LASSO 患者分数预测（不同任务，不能与 LOCO AUROC 直接比较）",
        "",
        patient[["state_id", "n_fold_units", "median_test_r2", "median_test_pearson", "median_test_spearman", "median_test_high_low_auroc"]].to_markdown(index=False, floatfmt=".4f"),
        "",
        "## 同一患者折、同一 held-out association label 的匹配比较",
        "",
        selected_pf.to_markdown(index=False, floatfmt=".4f"),
        "",
        "## 解释",
        "",
        "- LASSO 的优势：直接、可解释，擅长在同癌种内从表达预测 RNAss/DNAss 连续值。RNAss 本身来源于转录组，因此这一任务尤其容易。",
        "- CC-HHGT 的优势：目标是 lncRNA–state 关联重现，不是患者分数回归；可以整合图拓扑、方向、患者效应、多 seed 不确定性，并在 LOCO 中评分完全未见癌种。",
        "- 只有 matched PF 表可以回答‘哪个方法更会找到可重现的 lncRNA–RNAss/DNAss 关系’；患者高低分 AUROC 回答的是另一个问题。",
        "",
        "## /public7 遗留结果的边界",
        "",
        "当前工作站未挂载 `${DATA_ROOT}/stem`，因此未把其中训练内图表数值当作正式基线。本报告用相同 TCGA stemness 分数、lncRNA 表达和既有患者 folds 重建了外层留出的 nested LASSO。若远端脚本/预测文件随后挂载，可追加逐文件一致性审计。",
        "",
    ]
    output.write_text("\n".join(lines), encoding="utf-8")


def build_ppt(
    matrix_audit: dict,
    task_metrics: pd.DataFrame,
    oof_summary: pd.DataFrame,
    matched_summary: pd.DataFrame,
    lasso_patient: pd.DataFrame,
    paired_tests: pd.DataFrame,
    figures: dict[str, Path],
    output: Path,
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    add_text(slide, "CC-HHGT V3.0", 0.7, 1.2, 12.0, 0.7, 28, True, COLORS["blue"])
    add_text(slide, "RNAss / DNAss 完整重训与 LASSO 公平对照", 0.7, 2.05, 12.0, 1.1, 30, True)
    add_text(slide, "LOCO 跨癌泛化 × PF 患者间关联重现 × nested LASSO 患者分数预测", 0.75, 3.35, 11.8, 0.65, 17, False, COLORS["gray"])

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "1. 完整重训审计通过", "旧版全正样本训练矩阵被完全隔离")
    positive_rate = task_metrics.state_training_positive_rate.astype(float)
    add_bullets(
        slide,
        [
            f"任务矩阵：{matrix_audit['valid_task_metrics']}/{matrix_audit['expected_tasks']}（3 模型 × 33 癌种 × 3 seeds）",
            "每任务 state 训练集固定为 150,000 条：30,000 positive + 120,000 unlabeled",
            f"正样本率范围 {positive_rate.min():.3f}–{positive_rate.max():.3f}；缺类、缺校准、缺文件均 fail-closed",
            "OOF 发布前再次要求每模型恰好 99 个 calibrated runs",
        ],
        0.8,
        1.55,
        11.8,
        4.8,
        19,
    )

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "2. LOCO：完全未见癌种的 RNAss / DNAss 泛化")
    add_picture(slide, figures["loco"], 0.55, 1.25, 12.2)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "3. 为什么 LASSO 会显得非常强？", "它回答患者分数预测，而不是 lncRNA 关联重现")
    add_picture(slide, figures["endpoint"], 1.15, 1.28, 10.9)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "4. 公平比较：相同 patient fold、相同 held-out association label")
    add_picture(slide, figures["pf"], 0.75, 1.25, 11.8)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "5. 相对 LASSO 系数的配对增益")
    add_picture(slide, figures["delta"], 1.15, 1.28, 10.9)

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "6. 跨癌异质性仍是主要限制")
    add_picture(slide, figures["heatmap"], 0.55, 1.25, 5.2)
    add_bullets(
        slide,
        [
            "完整重训修复了优化目标，但不会消除癌种间真实生物学异质性。",
            "RNAss/DNAss 若仍低于 0.8，不代表训练没跑够；需要癌种条件化或 state-specific decoder。",
            "因此报告同时给 pooled、macro/median cancer 和逐癌种 AUROC。",
        ],
        6.15,
        1.7,
        6.3,
        3.8,
        18,
    )

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "7. 两类方法各自真正的优势")
    add_bullets(
        slide,
        [
            "LASSO：同癌种患者分数预测强、线性系数直观、部署简单。",
            "CC-HHGT/PF：更适合寻找能在未见患者中重现的 lncRNA–state 关系。",
            "CC-HHGT/LOCO：无需目标癌种 RNAss/DNAss 标签即可评分新癌种；LASSO 在该设定下无法拟合。",
            "图模型还提供方向、多模型一致性、embedding 覆盖和多 seed 不确定性。",
        ],
        0.8,
        1.55,
        11.8,
        4.8,
        19,
    )

    slide = presentation.slides.add_slide(blank)
    add_title(slide, "8. 结论与证据边界")
    add_bullets(
        slide,
        [
            "不要用 LASSO 患者高低分 AUROC 与 CC-HHGT 关联 AUROC 直接排高低。",
            "正式优势以 matched PF 配对表和 LOCO OOF 为证据；旧全正样本结果全部作废。",
            "当前工作站未挂载 /public7；已重建无泄漏 nested LASSO，远端原脚本挂载后可追加逐文件复核。",
        ],
        0.85,
        1.65,
        11.7,
        4.4,
        19,
    )
    presentation.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build final fixed-sampler RNAss/DNAss versus LASSO report")
    parser.add_argument("--matrix-audit-dir", required=True)
    parser.add_argument("--oof-analysis-dir", required=True)
    parser.add_argument("--matched-comparison-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    matrix_dir = Path(args.matrix_audit_dir).resolve()
    oof_dir = Path(args.oof_analysis_dir).resolve()
    matched_dir = Path(args.matched_comparison_dir).resolve()
    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    matrix_audit = require_pass(matrix_dir / "retrained_matrix_audit.json", "retrained matrix")
    require_pass(oof_dir / "state_oof_analysis_audit.json", "strict state OOF")
    require_pass(matched_dir / "MATCHED_LASSO_GRAPH_AUDIT.json", "matched LASSO/graph")
    task_metrics = pd.read_csv(matrix_dir / "retrained_task_metrics.tsv", sep="\t")
    oof_summary = pd.read_csv(oof_dir / "state_oof_state_summary.tsv", sep="\t")
    cancer_metrics = pd.read_csv(oof_dir / "state_oof_cancer_state_metrics.tsv", sep="\t")
    matched_summary = pd.read_csv(matched_dir / "matched_patient_fold_association_summary.tsv", sep="\t")
    lasso_patient = pd.read_csv(matched_dir / "lasso_patient_score_summary.tsv", sep="\t")
    paired_tests = pd.read_csv(matched_dir / "paired_auroc_tests_vs_lasso.tsv", sep="\t")
    for state_id in TARGETS:
        if state_id not in set(oof_summary.state_id) or state_id not in set(matched_summary.state_id):
            raise RuntimeError(f"Missing formal target in report inputs: {state_id}")

    configure_plotting()
    figures = {
        "loco": output_dir / "01_loco_rnass_dnass.png",
        "pf": output_dir / "02_matched_pf_vs_lasso.png",
        "endpoint": output_dir / "03_lasso_endpoint_contrast.png",
        "delta": output_dir / "04_paired_delta_vs_lasso.png",
        "heatmap": output_dir / "05_loco_cancer_heatmap.png",
    }
    save_loco_plot(oof_summary, figures["loco"])
    save_pf_plot(matched_summary, figures["pf"])
    save_endpoint_plot(lasso_patient, matched_summary, figures["endpoint"])
    save_delta_plot(paired_tests, figures["delta"])
    save_cancer_heatmap(cancer_metrics, figures["heatmap"])
    build_markdown(
        matrix_audit,
        task_metrics,
        oof_summary,
        matched_summary,
        lasso_patient,
        paired_tests,
        output_dir / "CC_HHGT_v3_RNAss_DNAss_vs_LASSO_report.md",
    )
    build_ppt(
        matrix_audit,
        task_metrics,
        oof_summary,
        matched_summary,
        lasso_patient,
        paired_tests,
        figures,
        output_dir / "CC_HHGT_v3_RNAss_DNAss_vs_LASSO_report.pptx",
    )
    payload = {
        "status": "PASS",
        "matrix_tasks": int(matrix_audit["valid_task_metrics"]),
        "figures": {key: str(path) for key, path in figures.items()},
        "markdown": str(output_dir / "CC_HHGT_v3_RNAss_DNAss_vs_LASSO_report.md"),
        "pptx": str(output_dir / "CC_HHGT_v3_RNAss_DNAss_vs_LASSO_report.pptx"),
    }
    (output_dir / "REPORT_SUCCESS.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
