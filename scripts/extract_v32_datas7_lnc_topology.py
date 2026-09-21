#!/usr/bin/env python3
"""Extract static lncRNA peak topology from official TCGA ATAC Data S7.

Only peak name, linked-gene symbol and linked-gene type are consumed.  The
workbook's association, correlation and FDR values are deliberately excluded
so labels/outcomes cannot leak into the static graph topology.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import json
import os
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def candidate_ids(path: Path) -> set[str]:
    with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if reader.fieldnames is None or "lncrna_id" not in reader.fieldnames:
            raise RuntimeError("Candidate key table lacks lncrna_id")
        return {row["lncrna_id"].removeprefix("LNC:").split(".")[0] for row in reader}


def promoter_symbol_map(path: Path, allowed: set[str]) -> tuple[dict[str, str], dict[str, list[str]]]:
    symbols: dict[str, set[str]] = defaultdict(set)
    with gzip.open(path, "rt", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if reader.fieldnames is None or not {"gene_id", "gene_name"}.issubset(reader.fieldnames):
            raise RuntimeError("GENCODE promoter table lacks gene_id/gene_name")
        for row in reader:
            gene_id = row["gene_id"].split(".")[0]
            if gene_id in allowed and row["gene_name"]:
                symbols[row["gene_name"]].add(gene_id)
    unique = {symbol: next(iter(ids)) for symbol, ids in symbols.items() if len(ids) == 1}
    ambiguous = {symbol: sorted(ids) for symbol, ids in symbols.items() if len(ids) > 1}
    return unique, ambiguous


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--promoters", type=Path, required=True)
    parser.add_argument("--candidate-keys", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    args = parser.parse_args()

    allowed = candidate_ids(args.candidate_keys)
    symbol_map, ambiguous = promoter_symbol_map(args.promoters, allowed)
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    worksheet = workbook["All_Links"]
    header = None
    header_row = None
    for index, values in enumerate(worksheet.iter_rows(values_only=True), 1):
        if values and values[0] == "Chromosome" and "Peak_Name" in values and "Linked_Gene" in values:
            header = list(values)
            header_row = index
            break
    if header is None or header_row is None:
        raise RuntimeError("Could not locate All_Links header in Data S7")
    peak_index = header.index("Peak_Name")
    gene_index = header.index("Linked_Gene")
    type_index = header.index("Linked_Gene_Type")
    links: set[tuple[str, str, str, str]] = set()
    total_rows = 0
    mapped_rows = 0
    for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        total_rows += 1
        peak = values[peak_index] if peak_index < len(values) else None
        symbol = values[gene_index] if gene_index < len(values) else None
        gene_type = values[type_index] if type_index < len(values) else None
        if not peak or not symbol or str(symbol) not in symbol_map:
            continue
        mapped_rows += 1
        gene_id = symbol_map[str(symbol)]
        links.add((f"LNC:{gene_id}", str(peak), str(symbol), str(gene_type or "")))
    rows = sorted(links)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.output.exists() or args.audit.exists():
        raise FileExistsError("DataS7 topology output reuse is forbidden")
    partial = args.output.with_name(args.output.name + ".partial")
    with gzip.open(partial, "wt", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(["lncrna_id", "peak_name", "linked_gene_symbol", "linked_gene_type", "source_sheet"])
        for lnc, peak, symbol, gene_type in rows:
            writer.writerow([lnc, peak, symbol, gene_type, "All_Links"])
    if partial.stat().st_size <= 0:
        raise RuntimeError("Data S7 topology extract is empty")
    os.replace(partial, args.output)
    audit = {
        "format": "CC_HHGT_V3_2_DATAS7_STATIC_LNCRNA_TOPOLOGY_V1",
        "status": "PASS_DATAS7_STATIC_TOPOLOGY",
        "source_workbook": str(args.workbook),
        "source_workbook_sha256": sha256(args.workbook),
        "source_sheet": "All_Links",
        "columns_consumed": ["Peak_Name", "Linked_Gene", "Linked_Gene_Type"],
        "correlation_columns_consumed": [],
        "fdr_columns_consumed": [],
        "candidate_gene_ids": len(allowed),
        "unique_symbol_mappings": len(symbol_map),
        "ambiguous_symbol_mappings_excluded": len(ambiguous),
        "source_data_rows": total_rows,
        "mapped_source_rows": mapped_rows,
        "unique_lnc_peak_links": len(rows),
        "output": str(args.output),
        "output_sha256": sha256(args.output),
    }
    audit_partial = args.audit.with_name(args.audit.name + ".partial")
    audit_partial.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(audit_partial, args.audit)
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    main()
